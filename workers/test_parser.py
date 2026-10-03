import tempfile
import unittest
from pathlib import Path
from parse_document import parse_document

def write_native_pdf(path):
    """Create an original text PDF fixture without a downloaded paper or extra dependency."""
    content = b'BT /F1 20 Tf 60 730 Td (Origin research) Tj 0 -40 Td (Table: First release) Tj ET'
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length ' + str(len(content)).encode() + b' >>\nstream\n' + content + b'\nendstream',
    ]
    pdf = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f'{number} 0 obj\n'.encode() + body + b'\nendobj\n')
    xref = len(pdf)
    pdf.extend(b'xref\n0 6\n0000000000 65535 f \n')
    for offset in offsets[1:]:
        pdf.extend(f'{offset:010} 00000 n \n'.encode())
    pdf.extend(f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    path.write_bytes(pdf)
    return path

class NormalParser(unittest.TestCase):
    def test_mixed_native_and_scanned_pdf_retains_both_pages(self):
        from PIL import Image, ImageDraw, ImageFont
        import pypdfium2 as pdfium
        image=Image.new('RGB',(1000,200),'white'); ImageDraw.Draw(image).text((40,50),'SCANNED ORIGIN PAGE',font=ImageFont.truetype('/System/Library/Fonts/Helvetica.ttc',48),fill='black')
        scan=self.root/'scan.pdf';image.save(scan,'PDF'); target=self.root/'mixed.pdf'
        with pdfium.PdfDocument.new() as merged, pdfium.PdfDocument(write_native_pdf(self.root/'native.pdf')) as native, pdfium.PdfDocument(scan) as scanned:
            merged.import_pages(native);merged.import_pages(scanned);merged.save(target)
        result=parse_document(target)
        self.assertTrue(any(b.get('page')==1 for b in result['blocks']))
        self.assertTrue(any(b.get('page')==2 and 'SCANNED' in b['text'] for b in result['blocks']))
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def test_markdown_preserves_original_text(self):
        p = self.root / 'notes.md'; p.write_text('# 访谈\n\n用户希望可以保存原始想法。', encoding='utf-8')
        self.assertIn('用户希望可以保存原始想法', parse_document(p)['content'])
    def test_docx_retains_chinese_paragraph(self):
        from docx import Document
        d = Document(); d.add_heading('研究', 0); d.add_paragraph('资料应该留在本机'); p = self.root / 'research.docx'; d.save(p)
        result = parse_document(p)
        self.assertIn('资料应该留在本机', result['content'])
        self.assertTrue(result['blocks'])
    def test_native_pdf_has_page_provenance(self):
        p = write_native_pdf(self.root / 'native.pdf')
        result = parse_document(p)
        self.assertTrue(any(block.get('page') == 1 for block in result['blocks']))
        self.assertIn('Table', result['content'])
    def test_slides_retain_chinese_content(self):
        from pptx import Presentation
        p = self.root / 'design.pptx'; deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = '产品方案'; slide.placeholders[1].text = '保留来源'; deck.save(p)
        self.assertIn('保留来源', parse_document(p)['content'])
    def test_spreadsheet_retains_table_data(self):
        from openpyxl import Workbook
        p = self.root / 'research.xlsx'; book = Workbook(); book.active.append(['功能', '优先级']); book.active.append(['原始资料', '最高']); book.save(p)
        self.assertIn('原始资料', parse_document(p)['content'])
    def test_csv_retains_chinese_rows(self):
        p = self.root / 'data.csv'; p.write_text('功能,优先级\n引用,高', encoding='utf-8')
        self.assertIn('引用', parse_document(p)['content'])
    def test_image_ocr_reads_visible_text(self):
        from PIL import Image, ImageDraw, ImageFont
        p = self.root / 'note.png'; image = Image.new('RGB', (1000, 200), 'white'); draw = ImageDraw.Draw(image)
        draw.text((40, 50), 'ORIGIN WORKSPACE 2026', font=ImageFont.truetype('/System/Library/Fonts/Helvetica.ttc', 48), fill='black'); image.save(p)
        self.assertIn('ORIGIN', parse_document(p)['content'])

class AdversarialParser(unittest.TestCase):
    def test_empty_pdf_page_is_reported_as_partial_not_silently_complete(self):
        from PIL import Image
        p=self.root/'blank.pdf';Image.new('RGB',(500,500),'white').save(p,'PDF')
        result=parse_document(p)
        self.assertEqual(result['status'],'partial_success')
        self.assertIn('PAGE_1_NO_TEXT',result['errors'])
    def setUp(self): self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def test_corrupt_pdf_fails_without_fabricated_text(self):
        p = self.root / 'corrupt.pdf'; p.write_bytes(b'not a pdf')
        with self.assertRaisesRegex(ValueError, 'PARSE_FAILED'): parse_document(p)
    def test_executable_is_not_a_document(self):
        p = self.root / 'run.sh'; p.write_text('touch /tmp/should-not-exist')
        with self.assertRaisesRegex(ValueError, 'UNSUPPORTED_FORMAT'): parse_document(p)
    def test_html_does_not_read_local_image(self):
        p = self.root / 'page.html'; p.write_text('<h1>正常资料</h1><script>EVIL_SCRIPT</script><img src="file:///etc/passwd"><p>保留正文</p>')
        result = parse_document(p)
        self.assertIn('保留正文', result['content'])
        self.assertNotIn('EVIL_SCRIPT', result['content'])
    def test_missing_file_has_fixed_error(self):
        with self.assertRaisesRegex(ValueError, 'FILE_NOT_FOUND'): parse_document(self.root / 'absent.md')
    def test_many_archive_entries_rejected_before_parsing(self):
        import zipfile
        p = self.root / 'crowded.docx'
        with zipfile.ZipFile(p, 'w') as z:
            for i in range(10_001): z.writestr(str(i), '')
        with self.assertRaisesRegex(ValueError, 'ARCHIVE_TOO_LARGE'): parse_document(p)
    def test_oversized_sparse_file_rejected(self):
        p = self.root / 'huge.pdf'
        with p.open('wb') as f: f.seek(100 * 1024 * 1024); f.write(b'x')
        with self.assertRaisesRegex(ValueError, 'FILE_TOO_LARGE'): parse_document(p)
    def test_invalid_utf8_is_an_explicit_failure(self):
        p = self.root / 'wrong.txt'; p.write_bytes(b'\xff\xfe\x00')
        with self.assertRaisesRegex(ValueError, 'PARSE_FAILED'): parse_document(p)
    def test_blank_image_is_partial_not_success(self):
        from PIL import Image
        p = self.root / 'blank.png'; Image.new('RGB', (100, 100), 'white').save(p)
        self.assertEqual(parse_document(p)['status'], 'partial_success')

if __name__ == '__main__': unittest.main()
